from __future__ import annotations

import base64
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Literal

from docx import Document
from openpyxl import load_workbook


ProcessedFileKind = Literal["text", "content_blocks"]

SUPPORTED_EXTENSIONS = "pdf, png, jpg, jpeg, gif, webp, docx, xlsx"
IMAGE_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}


class UnsupportedFileTypeError(ValueError):
    pass


@dataclass(frozen=True)
class ProcessedFile:
    kind: ProcessedFileKind
    content: str | list[dict[str, object]]


def process_file(
    filename: str,
    file_bytes: bytes,
    content_type: str | None = None,
) -> ProcessedFile:
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return ProcessedFile(
            kind="content_blocks",
            content=[_base64_block("document", "application/pdf", file_bytes)],
        )

    if extension in IMAGE_MEDIA_TYPES:
        media_type = (
            content_type
            if content_type in IMAGE_MEDIA_TYPES.values()
            else IMAGE_MEDIA_TYPES[extension]
        )
        return ProcessedFile(
            kind="content_blocks",
            content=[_base64_block("image", media_type, file_bytes)],
        )

    if extension == ".docx":
        return ProcessedFile(kind="text", content=_extract_docx_text(file_bytes))

    if extension == ".xlsx":
        return ProcessedFile(kind="text", content=_extract_xlsx_text(file_bytes))

    label = extension or Path(filename).name or "unknown"
    raise UnsupportedFileTypeError(
        f"지원하지 않는 파일 형식입니다: {label}. " f"지원 형식: {SUPPORTED_EXTENSIONS}"
    )


def _base64_block(
    block_type: str,
    media_type: str,
    file_bytes: bytes,
) -> dict[str, object]:
    return {
        "type": block_type,
        "source": {
            "type": "base64",
            "media_type": media_type,
            "data": base64.b64encode(file_bytes).decode("utf-8"),
        },
    }


def _extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                lines.append("\t".join(values))

    return "\n".join(lines)


def _extract_xlsx_text(file_bytes: bytes) -> str:
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    sheets: list[str] = []

    for worksheet in workbook.worksheets:
        lines = [f"[{worksheet.title}]"]
        for row in worksheet.iter_rows(values_only=True):
            values = [_format_cell_value(value) for value in row if value is not None]
            if values:
                lines.append("\t".join(values))
        sheets.append("\n".join(lines))

    workbook.close()
    return "\n\n".join(sheets)


def _format_cell_value(value) -> str:
    return str(value).strip()
