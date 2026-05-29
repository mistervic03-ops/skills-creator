import os
import sys
from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from services.file_processor import (
    UnsupportedFileTypeError,
    process_file,
)


def test_docx_text_extraction():
    document = Document()
    document.add_paragraph("첫 번째 문단")
    document.add_paragraph("두 번째 문단")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "표 A"
    table.rows[0].cells[1].text = "표 B"

    buffer = BytesIO()
    document.save(buffer)

    result = process_file("sample.docx", buffer.getvalue())

    assert result.kind == "text"
    assert result.content == "첫 번째 문단\n두 번째 문단\n표 A\t표 B"


def test_xlsx_text_extraction():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "요약"
    sheet.append(["항목", "내용"])
    sheet.append(["목적", "회의록 정리"])
    workbook.create_sheet("빈시트")

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    result = process_file("sample.xlsx", buffer.getvalue())

    assert result.kind == "text"
    assert result.content == ("[요약]\n항목\t내용\n목적\t회의록 정리\n\n[빈시트]")


def test_unsupported_file_type_error_message():
    with pytest.raises(UnsupportedFileTypeError) as exc_info:
        process_file("notes.txt", b"hello")

    assert str(exc_info.value) == (
        "지원하지 않는 파일 형식입니다: .txt. "
        "지원 형식: pdf, png, jpg, jpeg, gif, webp, docx, xlsx"
    )
